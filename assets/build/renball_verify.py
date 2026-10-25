"""
renball_verify.py — every figure traces to data.json
====================================================

Two checks, one command. Both run per project via a thin code/verify.py.

1. NUMBERS (fails on any mismatch)
   Every number printed in the project's copy must equal a numeric value in
   data/data.json at the precision shown — or that value x100 when a
   proportion is shown as a percent — or be covered by an explicit allow-rule
   that carries a reason. Copy sources:

     analysis.html   DOM text *after* JS runs (headless Chrome), so numbers
                     formatted in JS are checked too
     dashboard.html  static help text (its tables come from dashboard JSON)
     methodology.md  prose + tables (code spans and blocks skipped)
     slides.pptx     every text frame, table cell and group
     social/         thread, posts, alt hooks, captions, alt text, guide
     data.json       sentence-like strings (takeaways, findings)
     config.json     title / excerpt / finding

   Derived figures (ratios, totals, differences) must be written into
   data.json by build.py — the page may not compute a number the verifier
   can't trace. Bare integers below 13 and years are ignored (counts of
   leagues, seasons, slides); decimals, percentages and pp never are.
   An explicit sign (− or +) must agree with the sign in data.json; a bare
   whole number (a count) only matches whole numbers, never a rounded rate.

   ANCHORING. Equal to *some* value is not enough: in a large data.json
   almost any 2-decimal number matches something. Every match must also be
   anchored to its source, else it fails as WEAK:
     unique      only one data.json value shows as this number (the same
                 figure stored in several fields counts once)
     co-located  a candidate shares its parent or grandparent object with a
                 precisely identified number in the same unit of copy —
                 "60/79 = 75.95%" all sit under primary.treatment
     bound       an ANCHORS rule in code/verify.py pins the number to a
                 data.json path: (r"late go-ahead goal.*?\\d\\.\\d+×",
                 "$.early_vs_late_0to1.ratio"); its value is checked there.
                 On a page, data-verify-path="$.path" on the element does
                 the same for every number inside it (stat boxes).
   Units of copy: a paragraph / list item / table row (HTML), a line
   (Markdown), a slide, an X post, a config field.

2. STANDARD (errors fail, warnings print)
   Deliverables present, fresh (PDF newer than its source), consistent
   (config.json == projects.json entry, has_dashboard matches reality),
   takeaways well-formed, X package valid.

Per-project wrapper (code/verify.py):

    import renball_verify as V
    ALLOW = [
        (r"\\b85th minute", "illustrative minute in the hook, not a data claim"),
    ]
    ANCHORS = [
        (r"MDE\\D{0,20}\\d+\\.\\d", "$.primary.test.mde_pp"),
    ]
    sys.exit(V.main(PROJECT_DIR, allow=ALLOW, anchors=ANCHORS))

CLI:
    python code/verify.py              # both checks
    python code/verify.py --report     # also list every matched number, its path + anchor
    python code/verify.py --static     # skip headless render (raw HTML text)
"""

from __future__ import annotations

import argparse
import bisect
import functools
import html.parser
import json
import math
import re
import sys
import threading
from dataclasses import dataclass
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Iterable, Optional

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(REPO_ROOT / "assets" / "social"))

MIN_BARE_INT = 13

# Spans blanked before tokenising: never data claims
_BLANK = [re.compile(p) for p in (
    r"\b(?:19|20)\d{2}\s*[-–/]\s*(?:19|20)?\d{2}\b",           # seasons 2017-2018, 2023-24
    r"\b\d{4}-\d{2}-\d{2}(?:[T ][\d:.]+Z?)?",                    # ISO dates / timestamps
    r"\b\d{1,2}:\d{2}(?::\d{2})?\b",                             # clock times
    r"https?://\S+",
    r"\b[\w.-]+\.(?:com|org|net|io|html|json|md|py|mjs|js|css|pdf|png|pptx|csv|pkl|txt)\b(?:/\S*)?",
    r"(?:\.\./|\./)?(?:projects|assets|code|data|social|images)/\S+",   # repo paths
    r"#[0-9a-fA-F]{3,8}\b",                                      # hex colours
    r"\[\^?\d+\]",                                               # footnote refs
    r"\bv?\d+\.\d+\.\d+\b",                                      # versions
)]

_NUM = re.compile(
    r"(?<![\w.,])(?P<sign>[-+−–])?"
    r"(?:(?P<int>\d{1,3}(?:,\d{3})+(?!\d)|\d+)(?P<frac>\.\d+)?|(?P<bare>\.\d+))"
    r"(?P<suf>\s?%|\s?pp\b|\s?p\.p\.|×|x(?!\w)|k(?!\w))?(?!\w)")

DEFAULT_ALLOW: list[tuple[str, str]] = [
    (r"\b(?:80|90|95|99)\s?%\s*(?:CIs?\b|C\.I\.|confidence|credible|Wilson|Newcombe|bootstrap|HDI|HPD|"
     r"prediction|uncertainty|intervals?)", "interval level"),
    (r"(?:CI|confidence|credible|interval)s?\s*(?:level\s*)?(?:of\s*|at\s*)?\(?(?:80|90|95|99)\s?%",
     "interval level"),
    (r"(?:α|alpha)\s*(?:=|of)?\s*0?\.\d+", "significance level"),
    (r"\bp\s*[<≤>≥]\s*0?\.\d+", "p-value threshold"),
    (r"\b(?:80|90)\s?%\s*power|\bpower\s*(?:of\s*)?(?:80|90)\s?%", "power target"),
    (r"\b[Pp]roject\s*\d+\b|RESEARCH\s*·\s*\d+", "project number"),
    (r"(?m)^\s*\d{1,3}\s*/\s*\d{1,3}\s*$", "page / post counter"),
    (r"\b(?:1600|1080)\s*[x×]\s*(?:900|1350)\b", "image size"),
    (r"\bPython\s*\d+(?:\.\d+)*", "software version"),
    (r"\b\d+\s*/\s*280\b|\b280\s*(?:characters|chars)\b", "X character limit"),
    (r"\b90\s*(?:minutes|mins?)\b|\b90\+", "match length"),
    (r"(?:\bper[- ]?|/\s?)90\b", "per-90 unit"),
    (r"(?<![\d.,])(?:0|15|30|45|60|75)\s?[-–]\s?(?:15|30|45|60|75|90)\b(?!\s?%)", "15-minute bin label"),
    (r"§\s?\d+(?:\.\d+)*", "section reference"),
    (r"\b(?:top|Top|TOP)[ -]?5\b", "league scope"),
]


# ════════════════════════════════════════════════════════════════════════
# Truth set
# ════════════════════════════════════════════════════════════════════════
_PATH_PARTS = re.compile(r"\.[^.\[\]]+|\[\d+\]")


class Truth:
    """Numeric leaves of data.json, searchable at any display precision."""

    def __init__(self, data):
        self.data = data
        self.paths: list[str] = []
        self.values: list[float] = []
        for path, v in self._leaves(data, "$"):
            self.paths.append(path)
            self.values.append(v)
        order = sorted(range(len(self.values)), key=lambda i: abs(self.values[i]))
        self._abs = [abs(self.values[i]) for i in order]
        self._idx = order

    def _range(self, x: float, tol: float) -> list[int]:
        lo = bisect.bisect_left(self._abs, x - tol)
        hi = bisect.bisect_right(self._abs, x + tol)
        return self._idx[lo:hi]

    def find_all(self, x: float, tol: float, *, pct: bool = False, sign: str = "",
                 whole: bool = False) -> list[int]:
        """Leaves that show as x at this precision. A % / pp token also matches a
        proportion (|v| <= 1) times 100; an explicit − / + must agree in sign.
        A bare whole number (a count) only matches whole numbers; a decimal or a
        percentage matches a count only when no measured (float) value fits."""
        hits = set(self._range(x, tol))
        if pct:
            hits |= {i for i in self._range(x / 100, tol / 100) if abs(self.values[i]) <= 1}
        hits = [i for i in sorted(hits) if _sign_ok(self.values[i], sign)]
        if whole:
            return [i for i in hits if float(self.values[i]).is_integer()]
        floats = [i for i in hits if isinstance(self.values[i], float)]
        return floats or hits

    def resolve(self, path: str):
        node = self.data
        for part in _PATH_PARTS.findall(path[1:] if path.startswith("$") else path):
            node = node[int(part[1:-1])] if part.startswith("[") else node[part[1:]]
        return node

    @staticmethod
    def scopes_of(path: str) -> set[str]:
        """Parent and grandparent objects of a leaf (never the root)."""
        parts = _PATH_PARTS.findall(path[1:])
        return {"$" + "".join(parts[:len(parts) - cut]) for cut in (1, 2) if len(parts) - cut >= 1}

    def scopes(self, i: int) -> set[str]:
        return self.scopes_of(self.paths[i])

    def identities(self, idx: Iterable[int]) -> set[float]:
        """Distinct values — one figure stored in several fields counts once."""
        return {round(float(self.values[i]), 9) for i in idx}

    @staticmethod
    def _leaves(node, path):
        if isinstance(node, bool) or node is None:
            return
        if isinstance(node, (int, float)):
            if math.isfinite(node):
                yield path, node
        elif isinstance(node, dict):
            for k, v in node.items():
                yield from Truth._leaves(v, f"{path}.{k}")
        elif isinstance(node, list):
            for k, v in enumerate(node):
                yield from Truth._leaves(v, f"{path}[{k}]")

    def __len__(self) -> int:
        return len(self.paths)

    def nearest(self, x: float) -> tuple[float, str]:
        i = bisect.bisect_left(self._abs, x)
        cands = [j for j in (i - 1, i) if 0 <= j < len(self._abs)]
        j = min(cands, key=lambda k: abs(self._abs[k] - x))
        return self.values[self._idx[j]], self.paths[self._idx[j]]


def _sign_ok(v: float, sign: str) -> bool:
    return not sign or v == 0 or (v < 0) == (sign == "−")


# ════════════════════════════════════════════════════════════════════════
# Copy extraction
# ════════════════════════════════════════════════════════════════════════
class _TextParser(html.parser.HTMLParser):
    """Text per block (paragraph, list item, table row, heading, div) — the
    unit for anchoring. Inline tags (strong, span, a …) don't split a block."""
    SKIP = {"script", "style", "noscript", "svg", "template", "code"}
    VOID = {"br", "img", "meta", "link", "input", "hr", "wbr", "source"}
    BLOCK = {"p", "div", "li", "ul", "ol", "tr", "table", "thead", "tbody", "h1", "h2", "h3", "h4", "h5",
             "h6", "section", "header", "footer", "article", "aside", "nav", "main", "blockquote", "figure",
             "figcaption", "dl", "dt", "dd", "pre", "body", "canvas", "details", "summary", "label", "button"}
    CELL = {"td", "th"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.depth_skip = 0
        self.stack: list[tuple[str, bool, Optional[str]]] = []
        self.chunks: list[tuple[str, Optional[str]]] = []
        self.buf: list[str] = []

    def bound(self) -> Optional[str]:
        return next((b for _, _, b in reversed(self.stack) if b), None)

    def flush(self):
        text = " ".join("".join(self.buf).split()).strip(" |")
        if text:
            self.chunks.append((text, self.bound()))
        self.buf = []

    def handle_starttag(self, tag, attrs):
        if tag in self.VOID:
            if tag == "br":
                self.buf.append(" ")
            return
        a = dict(attrs)
        bind = a.get("data-verify-path")
        if tag in self.BLOCK or bind:
            self.flush()
        elif tag in self.CELL:
            self.buf.append(" | ")
        skip = tag in self.SKIP or a.get("data-verify") == "skip"
        self.stack.append((tag, skip, bind))
        self.depth_skip += skip

    def handle_endtag(self, tag):
        while self.stack:
            t, skip, bind = self.stack[-1]
            if bind or (t == tag and tag in self.BLOCK):
                self.flush()                     # before the pop, so a bound unit keeps its path
            self.stack.pop()
            self.depth_skip -= skip
            if t == tag:
                break

    def handle_data(self, data):
        if not self.depth_skip and data.strip():
            self.buf.append(data)

    def close(self):
        super().close()
        self.flush()


def html_text(markup: str) -> list[tuple[str, Optional[str]]]:
    """(text, bound path) per block. An element with data-verify-path="$.path"
    is its own unit and every number in it is checked against that path."""
    p = _TextParser()
    p.feed(markup)
    p.close()
    return p.chunks


class _QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


@functools.lru_cache(maxsize=1)
def _server() -> str:
    handler = functools.partial(_QuietHandler, directory=str(REPO_ROOT))
    srv = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return f"http://127.0.0.1:{srv.server_address[1]}"


def rendered_html(page: Path) -> str:
    """
    DOM after JS (fetch(data.json), templating) via headless Chrome. The
    browser locale is pinned to en-GB (the site's number format); pages must
    pin it too — renball_qa.mjs catches pages that don't.
    """
    import renball_pdf
    url = f"{_server()}/{page.resolve().relative_to(REPO_ROOT).as_posix()}"
    res = renball_pdf.headless_chrome(["--lang=en-GB", "--accept-lang=en-GB",
                                       "--virtual-time-budget=12000", "--dump-dom", url], timeout=90)
    if "<body" not in res.stdout:
        raise RuntimeError(f"headless render of {page.name} failed: {res.stderr[-300:]}")
    return res.stdout


_MD_STRIP = [
    (re.compile(r"```.*?```", re.S), " "),
    (re.compile(r"`[^`\n]+`"), " "),
    (re.compile(r"<!--.*?-->", re.S), " "),
    (re.compile(r"\]\([^)]*\)"), "]"),
    (re.compile(r"(?m)^\s*(?:[-*+]\s+)?\d+[.)]\s+"), ""),           # list numbering
    (re.compile(r"(?m)^#+\s*(?:§\s*)?\d+(?:\.\d+)*\.?\s"), ""),       # numbered headings
]


def md_text(raw: str) -> str:
    """Copy text of a Markdown file. Stripped spans keep their line breaks, so
    line N of the result is line N of the file."""
    for rx, rep in _MD_STRIP:
        raw = rx.sub(lambda m: rep + "\n" * m.group(0).count("\n"), raw)
    return raw


def pptx_text(path: Path) -> list[tuple[str, str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield cell.text
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                yield sh.text_frame.text

    out = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [t for t in walk(slide.shapes) if t.strip()]
        if texts:
            out.append((f"slide {n}", "\n".join(texts)))      # one slide = one unit
    return out


def _sentences(node, path="$") -> Iterable[tuple[str, str]]:
    if isinstance(node, str):
        if len(node.split()) >= 6:
            yield path, node
    elif isinstance(node, dict):
        for k, v in node.items():
            yield from _sentences(v, f"{path}.{k}")
    elif isinstance(node, list):
        for i, v in enumerate(node):
            yield from _sentences(v, f"{path}[{i}]")


def collect_copy(project_dir: Path, data: dict, *, render: bool = True) -> list[tuple[str, str, str]]:
    """(source, location, text) for every piece of copy in the project. HTML
    units inside a data-verify-path element get the location "... @ $.path"."""
    pd = project_dir
    out: list[tuple[str, str, str]] = []

    def html_units(source: str, loc: str, markup: str) -> list[tuple[str, str, str]]:
        return [(source, f"{loc} @ {bind}" if bind else loc, t) for t, bind in html_text(markup)]

    page = pd / "analysis.html"
    if page.exists():
        markup = rendered_html(page) if render else page.read_text(encoding="utf-8")
        out += html_units("analysis.html", "rendered" if render else "static", markup)
    dash = pd / "dashboard.html"
    if dash.exists():
        out += html_units("dashboard.html", "static", dash.read_text(encoding="utf-8"))
    meth = pd / "methodology.md"
    if meth.exists():
        for i, line in enumerate(md_text(meth.read_text(encoding="utf-8")).splitlines(), start=1):
            if line.strip():
                out.append(("methodology.md", f"line {i}", line))
    slides = pd / "slides.pptx"
    if slides.exists():
        out += [("slides.pptx", loc, t) for loc, t in pptx_text(slides)]
    if (pd / "social").exists():
        import renball_social
        pkg = renball_social.Package.load(pd)
        out += [("social", loc, t) for loc, t in pkg.all_text()]
        guide = pd / "social" / "posting_guide.md"
        if guide.exists():
            for i, line in enumerate(md_text(guide.read_text(encoding="utf-8")).splitlines(), start=1):
                if line.strip():
                    out.append(("social", f"posting_guide.md line {i}", line))
    out += [("data.json", path, t) for path, t in _sentences(data)]
    cfg = pd / "config.json"
    if cfg.exists():
        c = json.loads(cfg.read_text(encoding="utf-8"))
        out += [("config.json", k, c[k]) for k in ("title", "excerpt", "finding") if c.get(k)]
    return out


# ════════════════════════════════════════════════════════════════════════
# Numbers check
# ════════════════════════════════════════════════════════════════════════
@dataclass
class Hit:
    source: str
    loc: str
    token: str
    value: float
    status: str             # bound | co-located | unique | allowed | weak | unmatched
    detail: str
    context: str
    cands: tuple = ()
    sign: str = ""


_ANY = re.compile(r"(?s).+")
ANCHORED = ("bound", "co-located", "unique")
STATUSES = (*ANCHORED, "allowed", "weak", "unmatched")
NARROW = 2          # a number with at most this many distinct candidates can anchor its neighbours


def _blank(text: str) -> str:
    for rx in _BLANK:
        text = rx.sub(lambda m: " " * len(m.group(0)), text)
    return text


def _shows_as(v, x: float, tol: float, pct: bool, sign: str) -> bool:
    if isinstance(v, bool) or not isinstance(v, (int, float)) or not _sign_ok(v, sign):
        return False
    return abs(abs(v) - x) <= tol or (pct and abs(v) <= 1 and abs(abs(v) * 100 - x) <= tol)


def scan(text: str, truth: Truth, allow: list[tuple[re.Pattern, str]],
         anchors: list[tuple[re.Pattern, str]] = (), *, source: str = "", loc: str = "") -> list[Hit]:
    """Every number in one unit of copy, matched to data.json and anchored."""
    t = _blank(text)
    allowed_spans = [(m.start(), m.end(), why) for rx, why in allow for m in rx.finditer(t)]
    bound_spans = [(m.start(), m.end(), path) for rx, path in anchors for m in rx.finditer(t)]
    hits: list[Hit] = []
    for m in _NUM.finditer(t):
        intpart, frac, bare, suf = m.group("int"), m.group("frac"), m.group("bare"), (m.group("suf") or "").strip()
        digits = (intpart or "0").replace(",", "")
        x = float(digits + (frac or bare or ""))
        dec = len((frac or bare or ".")[1:])
        if not frac and not bare and not suf:
            if x < MIN_BARE_INT and "," not in (intpart or ""):
                continue
            if 1900 <= x <= 2099 and len(digits) == 4 and "," not in intpart:
                continue
        tol = 0.5 * 10 ** -dec + 1e-9
        if suf == "k":
            x, tol = x * 1000, 500 * 10 ** -dec
        pct = suf in ("%", "pp", "p.p.")
        sign = {"-": "−", "−": "−", "+": "+"}.get(m.group("sign") or "", "")
        ctx = text[max(0, m.start() - 45):m.end() + 35].replace("\n", " ")
        h = Hit(source, loc, m.group(0).strip(), x, "", "", ctx, sign=sign)
        why = next((w for s, e, w in allowed_spans if s <= m.start() and m.end() <= e), None)
        bound = next((p for s, e, p in bound_spans if s <= m.start() < e), None)   # starts inside the span
        if why:
            h.status, h.detail = "allowed", why
        elif bound:
            try:
                v = truth.resolve(bound)
            except (KeyError, IndexError, TypeError):
                v = None
            if _shows_as(v, x, tol, pct, sign):
                h.status, h.detail = "bound", bound
            else:
                h.status, h.detail = "unmatched", f"ANCHORS binds it to {bound} = {v!r}"
        else:
            whole = not frac and not bare and not suf
            h.cands = tuple(truth.find_all(x, tol, pct=pct, sign=sign, whole=whole))
            if h.cands:
                h.status = "pending"
            else:
                near, npath = truth.nearest(x)
                h.status, h.detail = "unmatched", f"nearest {near:g} at {npath}"
        hits.append(h)

    # Anchoring: unique, or co-located with a precisely identified neighbour.
    def scopes(h: Hit) -> set[str]:
        if h.status == "bound":
            return Truth.scopes_of(h.detail)
        if h.status == "pending" and len(truth.identities(h.cands)) <= NARROW:
            return set().union(*(truth.scopes(i) for i in h.cands))
        return set()

    anchor_scopes = [scopes(h) for h in hits]
    for k, h in enumerate(hits):
        if h.status != "pending":
            continue
        others = set().union(*(s for j, s in enumerate(anchor_scopes) if j != k))
        near = [i for i in h.cands if truth.scopes(i) & others]
        n_id = len(truth.identities(h.cands))
        if near:
            h.status, h.detail = "co-located", truth.paths[near[0]]
        elif n_id == 1:
            h.status, h.detail = "unique", truth.paths[h.cands[0]]
        else:
            eg = ", ".join(truth.paths[i] for i in h.cands[:3])
            h.status, h.detail = "weak", f"{n_id} different data.json values show as {h.token}, e.g. {eg}"
    return hits


def check_numbers(project_dir: Path, *, allow: Iterable[tuple[str, str]] = (),
                  anchors: Iterable[tuple[str, str]] = (),
                  render: bool = True, report: bool = False) -> int:
    data = json.loads((project_dir / "data" / "data.json").read_text(encoding="utf-8"))
    truth = Truth(data)
    rules = [(re.compile(p, re.I), why) for p, why in [*DEFAULT_ALLOW, *allow]]
    binds = [(re.compile(p, re.I | re.S), path) for p, path in anchors]
    for _, path in binds:
        try:
            truth.resolve(path)
        except (KeyError, IndexError, TypeError):
            raise SystemExit(f"ANCHORS: {path} is not in data.json")
    hits: list[Hit] = []
    for source, loc, text in collect_copy(project_dir, data, render=render):
        unit_binds = binds
        if " @ $" in loc:                                   # data-verify-path element
            path = loc.split(" @ ", 1)[1]
            try:
                truth.resolve(path)
            except (KeyError, IndexError, TypeError):
                raise SystemExit(f"data-verify-path: {path} is not in data.json ({source})")
            unit_binds = [(_ANY, path), *binds]
        hits += scan(text, truth, rules, unit_binds, source=source, loc=loc)

    bad = [h for h in hits if h.status in ("weak", "unmatched")]
    print(f"== Numbers check: {project_dir.name} ==")
    print(f"  truth: {len(truth)} numeric values in data.json")
    by_src: dict[str, dict[str, int]] = {}
    for h in hits:
        c = by_src.setdefault(h.source, dict.fromkeys(STATUSES, 0))
        c[h.status] += 1
    for src, c in by_src.items():
        matched = sum(c[s] for s in ANCHORED)
        how = " · ".join(f"{c[s]} {s}" for s in ANCHORED if c[s])
        ok = not (c["weak"] or c["unmatched"])
        print(f"  [{'OK ' if ok else '!! '}] {src:<15} {matched:>4} matched ({how or 'none'})  "
              f"{c['allowed']:>3} allowed  {c['weak']:>3} weak  {c['unmatched']:>3} unmatched")
    if report:
        for h in hits:
            if h.status not in ("weak", "unmatched"):
                print(f"      {h.status:<10} {h.token:>12}  {h.source}:{h.loc}  <- {h.detail}")
    for h in bad:
        print(f"  [{'WEAK' if h.status == 'weak' else 'FAIL'}] {h.source} {h.loc}: '{h.token}'  ({h.detail})")
        print(f"         ...{h.context}...")
    if any(h.status == "weak" for h in bad):
        print("  WEAK = the number exists in data.json but can't be tied to one value. Put it next to its\n"
              "  sibling figures, or pin it with an ANCHORS rule in code/verify.py: (regex, '$.path').")
    return len(bad)


# ════════════════════════════════════════════════════════════════════════
# Standard check
# ════════════════════════════════════════════════════════════════════════
AUDIENCE_KEYS = ("coaches", "managers", "clubs", "agents")
EVIDENCE_KEYS = ("strong", "moderate", "exploratory")
REQUIRED_FILES = ("analysis.html", "config.json", "data/data.json", "methodology.md",
                  "methodology.html", "methodology.pdf", "slides.pptx", "slides.pdf")
FRESHNESS = (("slides.pdf", "slides.pptx"), ("methodology.html", "methodology.md"),
             ("methodology.pdf", "methodology.html"), ("slides.pptx", "data/data.json"))


def check_standard(project_dir: Path) -> int:
    pd = project_dir
    errors, warnings = [], []
    for f in REQUIRED_FILES:
        if not (pd / f).exists():
            errors.append(f"missing {f}")
    for newer, older in FRESHNESS:
        a, b = pd / newer, pd / older
        if a.exists() and b.exists() and a.stat().st_mtime < b.stat().st_mtime:
            warnings.append(f"{newer} is older than {older} -- rebuild/re-export it")

    data = json.loads((pd / "data" / "data.json").read_text(encoding="utf-8")) \
        if (pd / "data" / "data.json").exists() else {}
    tk = data.get("takeaways")
    if not isinstance(tk, dict):
        errors.append("data.json has no 'takeaways' object (four audiences)")
    else:
        for aud in AUDIENCE_KEYS:
            items = tk.get(aud)
            if not isinstance(items, list):
                errors.append(f"takeaways.{aud} missing (use [] for 'no defensible claim')")
                continue
            for i, it in enumerate(items):
                if it.get("evidence") not in EVIDENCE_KEYS:
                    errors.append(f"takeaways.{aud}[{i}].evidence = {it.get('evidence')!r}")
                if not it.get("text"):
                    errors.append(f"takeaways.{aud}[{i}] has no text")
                elif len(it["text"]) > 170:
                    warnings.append(f"takeaways.{aud}[{i}] is {len(it['text'])} chars "
                                    "(keep <= 140 so the slide card fits)")
            if len(items) > 2:
                warnings.append(f"takeaways.{aud} has {len(items)} items (slide fits 2)")

    cfg_path = pd / "config.json"
    if cfg_path.exists():
        cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
        if not cfg.get("finding"):
            errors.append("config.json has no 'finding'")
        entries = {p["folder"]: p for p in json.loads((REPO_ROOT / "projects.json").read_text(encoding="utf-8"))}
        if pd.name not in entries:
            errors.append("project missing from projects.json")
        elif entries[pd.name] != cfg:
            diff = sorted(k for k in set(cfg) | set(entries[pd.name]) if cfg.get(k) != entries[pd.name].get(k))
            errors.append(f"config.json != projects.json entry (fields: {', '.join(diff)})")
        if bool(cfg.get("has_dashboard")) != (pd / "dashboard.html").exists():
            errors.append("config.has_dashboard disagrees with dashboard.html on disk")

    page = pd / "analysis.html"
    if page.exists():
        src = page.read_text(encoding="utf-8")
        for needle, what in (("renball-components.js", "shared components script"),
                             ('data-renball="slide-viewer"', "embedded slide viewer"),
                             ("renderTakeaways", "takeaways grid"),
                             ("methodology.pdf", "methodology PDF link"),
                             ("slides.pdf", "slides PDF link")):
            if needle not in src:
                errors.append(f"analysis.html lacks the {what} ({needle})")

    if (pd / "social").exists():
        import renball_social
        errors += [f"X package: {e}" for e in renball_social.validate(
            renball_social.Package.load(pd), verbose=False)]
    else:
        errors.append("social/ X package missing")

    print(f"== Standard check: {pd.name} ==")
    for w in warnings:
        print(f"  [WARN] {w}")
    for e in errors:
        print(f"  [FAIL] {e}")
    if not errors and not warnings:
        print("  [OK ] all deliverables present, fresh and consistent")
    return len(errors)


def main(project_dir: Path, *, allow: Iterable[tuple[str, str]] = (),
         anchors: Iterable[tuple[str, str]] = (), argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=f"Verify {Path(project_dir).name}")
    ap.add_argument("--report", action="store_true", help="list matched numbers with their data.json paths")
    ap.add_argument("--static", action="store_true", help="skip the headless render of analysis.html")
    ap.add_argument("--numbers-only", action="store_true")
    ap.add_argument("--standard-only", action="store_true")
    a = ap.parse_args(argv)
    project_dir = Path(project_dir).resolve()
    fails = 0
    if not a.standard_only:
        fails += check_numbers(project_dir, allow=allow, anchors=anchors,
                               render=not a.static, report=a.report)
    if not a.numbers_only:
        fails += check_standard(project_dir)
    print("VERIFY PASSED" if not fails else f"VERIFY FAILED ({fails} problem(s))")
    return 1 if fails else 0


def selftest() -> int:
    """The matching rules on a toy data.json: python renball_verify.py --selftest"""
    data = {"primary": {"treatment": {"conversions": 60, "attempts": 79, "rate": 75.95},
                        "test": {"diff": -3.01, "mde_pp": 12.55}},
            "grid": [[{"v": 2.08}, {"v": 2.11}], [{"v": 2.14}, {"v": 0.7895}]],
            "ratio": 2.14, "rate_overall": 78.96}
    truth = Truth(data)
    rules = [(re.compile(p, re.I), w) for p, w in DEFAULT_ALLOW]
    binds = [(re.compile(r"late\D{0,40}\d\.\d+×", re.I), "$.ratio")]

    def statuses(text, anchors=()):
        return [h.status for h in scan(text, truth, rules, list(anchors))]

    cases = [
        ("60/79 = 75.95%", [], ["co-located", "co-located", "co-located"]),
        ("a −3.01 pp gap", [], ["unique"]),
        ("a +3.01 pp gap", [], ["unmatched"]),              # sign disagrees with data.json
        ("79 penalties", [], ["unique"]),                   # count, not the 78.96% rate
        ("about 79 percent", [], ["unique"]),
        ("78.96%", [], ["unique"]),
        ("78.95%", [], ["unique"]),                         # 0.7895 x 100
        ("2.1× more", [], ["weak"]),                        # 2.08 / 2.11 / 2.14 all show as 2.1
        ("a late winner is worth 2.1× more", binds, ["bound"]),
        ("a late winner is worth 2.2× more", binds, ["unmatched"]),
        ("between 60-75 and 75-90 minutes, EP per 90", [], ["allowed"] * 5),
        ("95% CI, see §3.4", [], ["allowed", "allowed"]),
    ]
    fails = 0
    for text, anchors, want in cases:
        got = statuses(text, anchors)
        ok = got == want
        fails += not ok
        print(f"  [{'OK ' if ok else 'FAIL'}] {text!r:46} {got}" + ("" if ok else f"  want {want}"))
    units = html_text('<div data-verify-path="$.primary.test.mde_pp"><b>12.6</b> pp</div>'
                      '<p>a <strong>60</strong>/79 = 75.95% b</p>')
    want_units = [("12.6 pp", "$.primary.test.mde_pp"), ("a 60/79 = 75.95% b", None)]
    fails += units != want_units
    print(f"  [{'OK ' if units == want_units else 'FAIL'}] html units + data-verify-path {units}")
    md = "a\n```\ncode\nblock\n```\nline 6 has 12.55\n"
    line = md_text(md).splitlines().index("line 6 has 12.55") + 1
    fails += line != 6
    print(f"  [{'OK ' if line == 6 else 'FAIL'}] methodology line numbers survive code blocks (line {line})")
    print("SELFTEST PASSED" if not fails else f"SELFTEST FAILED ({fails})")
    return 1 if fails else 0


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--selftest":
        sys.exit(selftest())
    if len(sys.argv) < 2:
        print("usage: python assets/build/renball_verify.py projects/XX-name [--report] [--static]\n"
              "       python assets/build/renball_verify.py --selftest")
        sys.exit(2)
    sys.exit(main(Path(sys.argv[1]), argv=sys.argv[2:]))
